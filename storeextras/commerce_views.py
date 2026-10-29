import csv
from decimal import Decimal
from django.contrib import messages
from django.contrib.admin.views.decorators import staff_member_required
from django.contrib.auth.decorators import login_required
from django.db.models import Sum, Count
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views.decorators.http import require_POST
from orders.models import Coupon, Order, OrderItem
from products.models import Product

@login_required
@require_POST
def apply_coupon(request):
    code = request.POST.get('coupon_code', '').strip().upper()
    coupon = Coupon.objects.filter(code=code, is_active=True).first()
    if not coupon:
        messages.error(request, 'That coupon code is not valid.')
        return redirect('checkout')
    now = timezone.now()
    if coupon.starts_at and coupon.starts_at > now or coupon.expires_at and coupon.expires_at < now:
        messages.error(request, 'This coupon is not currently active.')
        return redirect('checkout')
    if coupon.usage_limit and coupon.used_count >= coupon.usage_limit:
        messages.error(request, 'This coupon has reached its usage limit.')
        return redirect('checkout')
    from cart.models import Cart
    cart = Cart.objects.filter(user=request.user).first()
    subtotal = sum((item.total_price for item in cart.items.select_related('product', 'variant')), Decimal('0.00')) if cart else Decimal('0.00')
    if subtotal < coupon.minimum_order_amount:
        messages.error(request, f'The minimum order for this coupon is ₹{coupon.minimum_order_amount}.')
        return redirect('checkout')
    request.session['applied_coupon'] = coupon.code
    messages.success(request, f'Coupon {coupon.code} applied. Discount will be checked again at checkout.')
    return redirect('checkout')

@login_required
@require_POST
def remove_coupon(request):
    request.session.pop('applied_coupon', None)
    messages.success(request, 'Coupon removed.')
    return redirect('checkout')

@login_required
@require_POST
def cancel_order_request(request, order_id):
    order = get_object_or_404(Order, pk=order_id, user=request.user)
    if order.status not in [Order.STATUS_PENDING, Order.STATUS_CONFIRMED]:
        messages.error(request, 'This order can no longer be cancelled online. Please contact support.')
        return redirect('order_detail', order_id=order.id)
    reason = request.POST.get('reason', '').strip()[:255]
    order.cancellation_reason = reason or 'Customer requested cancellation'
    order.status = Order.STATUS_CANCELLED
    order.save(update_fields=['status', 'cancellation_reason'])
    for item in order.items.select_related('product'):
        if not item.is_cancelled:
            item.product.stock += item.quantity
            item.product.save(update_fields=['stock'])
            item.is_cancelled = True
            item.save(update_fields=['is_cancelled'])
    messages.success(request, 'Cancellation request processed. Any prepaid refund must be handled through the payment provider.')
    return redirect('order_detail', order_id=order.id)

@staff_member_required
def analytics(request):
    paid_orders = Order.objects.filter(payment_status='Paid').exclude(status=Order.STATUS_CANCELLED)
    context = {'orders_count': Order.objects.count(), 'paid_revenue': paid_orders.aggregate(total=Sum('total_amount'))['total'] or 0, 'customers_count': Order.objects.values('user_id').distinct().count(), 'low_stock': Product.objects.filter(is_available=True, stock__lte=5).order_by('stock')[:12], 'low_stock_count': Product.objects.filter(is_available=True, stock__lte=5).count(), 'top_products': OrderItem.objects.values('product__name').annotate(units=Sum('quantity'), revenue=Sum('price')).order_by('-units')[:10]}
    return render(request, 'dashboard/analytics.html', context)

@staff_member_required
def export_orders_csv(request):
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="orders-report.csv"'
    writer = csv.writer(response)
    writer.writerow(['Order ID', 'Customer', 'Email', 'Status', 'Payment status', 'Total', 'Created'])
    for order in Order.objects.select_related('user').order_by('-created_at'):
        writer.writerow([order.pk, order.name, order.email, order.status, order.payment_status, order.total_amount, order.created_at.isoformat()])
    return response

@staff_member_required
def export_stock_csv(request):
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="stock-report.csv"'
    writer = csv.writer(response)
    writer.writerow(['Product ID', 'Product', 'Category', 'Stock', 'Price', 'Available'])
    for product in Product.objects.select_related('category').order_by('name'):
        writer.writerow([product.pk, product.name, product.category.name, product.stock, product.price, product.is_available])
    return response


@staff_member_required
def export_sales_excel(request):
    from openpyxl import Workbook
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="sales-report.xlsx"'
    workbook = Workbook(); sheet = workbook.active; sheet.title = 'Sales'
    sheet.append(['Order ID', 'Customer', 'Email', 'Status', 'Payment Status', 'Total', 'Created'])
    for order in Order.objects.select_related('user').order_by('-created_at'):
        sheet.append([order.pk, order.name, order.email, order.status, order.payment_status, float(order.total_amount), order.created_at.replace(tzinfo=None)])
    workbook.save(response); return response

@staff_member_required
def export_inventory_excel(request):
    from openpyxl import Workbook
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="inventory-report.xlsx"'
    workbook = Workbook(); sheet = workbook.active; sheet.title = 'Inventory'
    sheet.append(['Product ID', 'Product', 'Category', 'Stock', 'Price', 'Available'])
    for product in Product.objects.select_related('category').order_by('name'):
        sheet.append([product.pk, product.name, product.category.name, product.stock, float(product.price), product.is_available])
    workbook.save(response); return response

@staff_member_required
def export_refunds_excel(request):
    from openpyxl import Workbook
    from orders.models import ReturnReplaceRequest
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="refunds-report.xlsx"'
    workbook = Workbook(); sheet = workbook.active; sheet.title = 'Returns and refunds'
    sheet.append(['Request ID', 'Order ID', 'Customer', 'Product', 'Type', 'Status', 'Refund status', 'Refund amount', 'Created'])
    rows = ReturnReplaceRequest.objects.select_related('order_item__order__user', 'order_item__product').order_by('-created_at')
    for item in rows:
        sheet.append([item.pk, item.order_item.order_id, item.order_item.order.user.username, item.order_item.product.name, item.request_type, item.status, item.refund_status, float(item.refund_amount), item.created_at.replace(tzinfo=None)])
    workbook.save(response); return response

@staff_member_required
def import_products_csv(request):
    if request.method != 'POST' or not request.FILES.get('csv_file'):
        messages.error(request, 'Choose a CSV file and submit it using the import form.')
        return redirect('dashboard:products')
    import io
    from products.models import Category
    try:
        decoded = request.FILES['csv_file'].read().decode('utf-8-sig')
        reader = csv.DictReader(io.StringIO(decoded))
        required = {'name', 'category', 'description', 'price', 'stock'}
        if not reader.fieldnames or not required.issubset(set(reader.fieldnames)):
            raise ValueError('CSV must contain name, category, description, price and stock columns.')
        created = 0
        for row in reader:
            name = (row.get('name') or '').strip()[:200]
            category_name = (row.get('category') or '').strip()[:100]
            if not name or not category_name:
                continue
            category, _ = Category.objects.get_or_create(name=category_name)
            Product.objects.create(name=name, category=category, description=(row.get('description') or '').strip(), price=Decimal(row['price']), stock=max(0, int(row['stock'])), is_available=True)
            created += 1
        messages.success(request, f'Imported {created} products. Review the catalog before publishing.')
    except (UnicodeDecodeError, ValueError, KeyError, TypeError, ArithmeticError) as exc:
        messages.error(request, f'Import stopped: {exc}')
    return redirect('dashboard:products')
